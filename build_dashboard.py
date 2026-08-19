import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

SRC = "Property_Management_Tracker.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=False)

if "Dashboard" in wb.sheetnames:
    del wb["Dashboard"]
ws = wb.create_sheet("Dashboard", 0)
ws.sheet_view.showGridLines = False

# ---- palette (matches existing workbook header color) ----
DARK_GREEN = "356854"
LIGHT_GREEN = "DCE6E0"
WHITE = "FFFFFF"
GREY_TEXT = "595959"
YELLOW = "FFFF00"
BLUE_FONT = "0000FF"

FONT_NAME = "Arial"

title_font = Font(name=FONT_NAME, size=18, bold=True, color=WHITE)
subtitle_font = Font(name=FONT_NAME, size=10, italic=True, color=WHITE)
section_font = Font(name=FONT_NAME, size=13, bold=True, color=WHITE)
subheader_font = Font(name=FONT_NAME, size=11, bold=True, color=DARK_GREEN)
colhead_font = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
label_font = Font(name=FONT_NAME, size=11, color="000000")
value_font = Font(name=FONT_NAME, size=11, bold=True, color="000000")
note_font = Font(name=FONT_NAME, size=9, italic=True, color=GREY_TEXT)
input_font = Font(name=FONT_NAME, size=11, bold=True, color=BLUE_FONT)
assumption_head_font = Font(name=FONT_NAME, size=11, bold=True, color="000000")

title_fill = PatternFill("solid", fgColor=DARK_GREEN)
section_fill = PatternFill("solid", fgColor=DARK_GREEN)
subheader_fill = PatternFill("solid", fgColor=LIGHT_GREEN)
colhead_fill = PatternFill("solid", fgColor="6E8F82")
input_fill = PatternFill("solid", fgColor=YELLOW)

thin = Side(style="thin", color="C9D3CF")
box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

CUR = '"$"#,##0.00'
INT = '#,##0'
PCT = '0.0%'
DATE_FMT = 'mmm d, yyyy'

# ---- headroom ranges (data + growth room), consistent everywhere below ----
PM = "'PM ONLY'"
PM_R = (2, 2000)
LO = "'LEASE ONLY'"
LO_R = (2, 1000)
LEN_ON = "'Len-Onboarded'"
LEN_R = (2, 500)

def col(sheet_ref, letter, r=None):
    start, end = r or (PM_R if sheet_ref == PM else LO_R if sheet_ref == LO else LEN_R)
    return f"{sheet_ref}!{letter}{start}:{letter}{end}"

def set_cell(coord, value=None, font=None, fill=None, fmt=None, align=None, border=None):
    c = ws[coord]
    if value is not None:
        c.value = value
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c

def merge_banner(row, text, font, fill, span="B:F", height=None):
    start, end = span.split(":")
    ws.merge_cells(f"{start}{row}:{end}{row}")
    set_cell(f"{start}{row}", text, font=font, fill=fill,
             align=Alignment(horizontal="left", vertical="center", indent=1))
    for cc in range(ord(start), ord(end) + 1):
        ws[f"{chr(cc)}{row}"].fill = fill
    if height:
        ws.row_dimensions[row].height = height

def metric_row(r, label, formula, fmt=INT, note=""):
    set_cell(f"B{r}", label, font=label_font)
    set_cell(f"C{r}", formula, font=value_font, fmt=fmt,
              align=Alignment(horizontal="right"))
    if note:
        set_cell(f"D{r}", note, font=note_font, align=Alignment(wrap_text=True, vertical="center"))
    ws[f"C{r}"].border = box_border

def table_header(r, headers):
    for i, h in enumerate(headers):
        cell = set_cell(f"{chr(ord('B')+i)}{r}", h, font=colhead_font, fill=colhead_fill,
                         align=Alignment(horizontal="center", vertical="center"))
    ws.row_dimensions[r].height = 18

# =========================================================
# Column widths
# =========================================================
widths = {"A": 2.5, "B": 40, "C": 16, "D": 55, "E": 16, "F": 16}
for cletter, w in widths.items():
    ws.column_dimensions[cletter].width = w

# =========================================================
# Title
# =========================================================
ws.row_dimensions[1].height = 34
merge_banner(1, "Property Management Tracker — Live Dashboard", title_font, title_fill)
ws.row_dimensions[2].height = 20
merge_banner(2, "Auto-calculated from PM ONLY, LEASE ONLY, and Len-Onboarded — all figures are live formulas that update when source rows change.",
             subtitle_font, title_fill)

set_cell("B4", "Report Date (auto)", font=assumption_head_font)
set_cell("C4", "=TODAY()", font=value_font, fmt=DATE_FMT, align=Alignment(horizontal="right"))
ws["C4"].border = box_border

set_cell("B5", "GST Rate (assumption, used for estimates below)", font=assumption_head_font)
set_cell("C5", 0.05, font=input_font, fill=input_fill, fmt=PCT, align=Alignment(horizontal="right"))
ws["C5"].border = box_border
set_cell("D5", "Input cell — edit if GST rate changes. PM ONLY's own 'TOTAL FEES WITH GST' column is blank in the source data, so this rate is used to estimate GST-inclusive management-fee revenue below.",
         font=note_font, align=Alignment(wrap_text=True, vertical="center"))

# =========================================================
# SECTION 1 — Portfolio Overview
# =========================================================
r = 7
merge_banner(r, "1.  Portfolio Overview", section_font, section_fill, height=22)
r += 1
table_header(r, ["Metric", "Value", "Notes"])
r += 1
row_total_units = r
metric_row(r, "Total Units Tracked (PM ONLY)", f"=COUNTA({col(PM,'D')})", INT,
           "Count of non-blank Property rows in PM ONLY.")
r += 1
row_active = r
metric_row(r, "Active Units", f'=COUNTIF({col(PM,"I")},"Active")', INT)
r += 1
row_terminated = r
metric_row(r, "Terminated Units", f'=COUNTIF({col(PM,"I")},"Terminated")', INT)
r += 1
row_occupied = r
metric_row(r, "Occupied Units", f'=COUNTIF({col(PM,"H")},"Occupied")', INT)
r += 1
row_vacant = r
metric_row(r, "Vacant Units", f'=COUNTIF({col(PM,"H")},"Vacant")', INT)
r += 1
row_takeover = r
metric_row(r, "Take Over (in transition) Units", f'=COUNTIF({col(PM,"H")},"Take over")', INT,
           "\"Take over\" is a third Occupancy Status value found in the source data, distinct from Occupied/Vacant.")
r += 1
row_occ_rate = r
metric_row(r, "Occupancy Rate (Occupied ÷ Occupied+Vacant+Take Over)",
           f"=IFERROR(C{row_occupied}/(C{row_occupied}+C{row_vacant}+C{row_takeover}),0)", PCT,
           "Guarded against a zero denominator.")
r += 2

# ---- Units by Assigned PM ----
merge_banner(r, "Units by Assigned PM (workload)", subheader_font, subheader_fill, height=18)
r += 1
table_header(r, ["Assigned PM", "Active Units", "Total Units (incl. Terminated)", "% of Active Portfolio"])
r += 1
pm_names = ["Thomas Bruchig", "Junny Ko", "Ranjeet Bhangu", "Hemi Persaud",
            "Lorraine Girouard", "Emman Lim", "Parmeet Singh", "Lorenzo Cabason", "Rachel Qin"]
pm_table_start = r
for name in pm_names:
    set_cell(f"B{r}", name, font=label_font)
    set_cell(f"C{r}", f'=COUNTIFS({col(PM,"F")},$B{r},{col(PM,"I")},"Active")',
              font=value_font, fmt=INT, align=Alignment(horizontal="right"), border=box_border)
    set_cell(f"D{r}", f'=COUNTIF({col(PM,"F")},$B{r})',
              font=value_font, fmt=INT, align=Alignment(horizontal="right"), border=box_border)
    set_cell(f"E{r}", f"=IFERROR(C{r}/$C${row_active},0)",
              font=value_font, fmt=PCT, align=Alignment(horizontal="right"), border=box_border)
    r += 1
pm_table_end = r - 1
r += 1

# =========================================================
# SECTION 2 — Revenue Snapshot
# =========================================================
merge_banner(r, "2.  Revenue Snapshot", section_font, section_fill, height=22)
r += 1
table_header(r, ["Metric", "Value", "Notes"])
r += 1
metric_row(r, "Total Monthly Rent Roll (PM ONLY)", f"=SUM({col(PM,'P')})", CUR)
r += 1
metric_row(r, "Total Monthly Rent Roll (LEASE ONLY, supplemental)", f"=SUM({col(LO,'M')})", CUR,
           "Leasing-only engagements; kept separate so it is not double-counted with PM ONLY rent.")
r += 1
row_mgmt_fee = r
metric_row(r, "Total Management Fee Revenue, excl. GST (PM ONLY)", f"=SUM({col(PM,'U')})", CUR)
r += 1
metric_row(r, "Total Management Fee Revenue, incl. GST (estimated)",
           f"=C{row_mgmt_fee}*(1+$C$5)", CUR,
           "Estimated using the GST Rate assumption above — PM ONLY's 'TOTAL FEES WITH GST' column has no source formula/values to sum directly.")
r += 1
metric_row(r, "Total Leasing Fee Revenue, excl. GST (PM ONLY)", f"=SUM({col(PM,'O')})", CUR)
r += 1
metric_row(r, "Total Leasing Fee Revenue, incl. GST (LEASE ONLY)", f"=SUM({col(LO,'Q')})", CUR,
           "LEASE ONLY column Q already computes Leasing Fee x 1.05 per row in the source.")
r += 1
row_disb_pm = r
metric_row(r, "Total Owner Disbursements (PM ONLY)", f"=SUM({col(PM,'X')})", CUR)
r += 1
row_disb_lo = r
metric_row(r, "Total Owner Disbursements (LEASE ONLY)", f"=SUM({col(LO,'S')})", CUR)
r += 1
metric_row(r, "Combined Total Owner Disbursements", f"=C{row_disb_pm}+C{row_disb_lo}", CUR)
r += 1
metric_row(r, "Ancillary Income — Pet Fee (PM ONLY)", f"=SUM({col(PM,'Q')})", CUR)
r += 1
metric_row(r, "Ancillary Income — Parking (PM ONLY)", f"=SUM({col(PM,'R')})", CUR)
r += 1
metric_row(r, "Ancillary Income — Storage (PM ONLY)", f"=SUM({col(PM,'S')})", CUR)
r += 2

# =========================================================
# SECTION 3 — Leasing Pipeline
# =========================================================
merge_banner(r, "3.  Leasing Pipeline", section_font, section_fill, height=22)
r += 1
table_header(r, ["Metric", "Value", "Notes"])
r += 1
row_expiring = {}
for days in (30, 60, 90):
    row_expiring[days] = r
    metric_row(r, f"Leases Expiring in Next {days} Days (Active units)",
               f'=COUNTIFS({col(PM,"I")},"Active",{col(PM,"N")},">="&$C$4,{col(PM,"N")},"<="&$C$4+{days})',
               INT, "Text/blank move-out dates are ignored automatically by COUNTIFS's numeric comparison.")
    r += 1
row_vacant_active = r
metric_row(r, "Vacant Units with Active PM Agreement (immediate revenue loss)",
           f'=COUNTIFS({col(PM,"H")},"Vacant",{col(PM,"I")},"Active")', INT)
r += 1
row_new_leases_month = r
metric_row(r, "New Leases Signed This Month (Lease Move In)",
           f'=COUNTIFS({col(PM,"M")},">="&DATE(YEAR($C$4),MONTH($C$4),1),{col(PM,"M")},"<="&$C$4)',
           INT, "Counts PM ONLY rows whose Lease Move In date falls within the current calendar month.")
r += 2

# =========================================================
# SECTION 4 — Growth
# =========================================================
merge_banner(r, "4.  Growth", section_font, section_fill, height=22)
r += 1
table_header(r, ["Metric", "Value", "Notes"])
r += 1
qstart = "DATE(YEAR($C$4),MONTH($C$4)-MOD(MONTH($C$4)-1,3),1)"
row_new_owners_month = r
metric_row(r, "New Owners Onboarded This Month (Len-Onboarded)",
           f'=COUNTIFS({col(LEN_ON,"A")},">="&DATE(YEAR($C$4),MONTH($C$4),1),{col(LEN_ON,"A")},"<="&$C$4)', INT)
r += 1
row_new_owners_quarter = r
metric_row(r, "New Owners Onboarded This Quarter (Len-Onboarded)",
           f'=COUNTIFS({col(LEN_ON,"A")},">="&{qstart},{col(LEN_ON,"A")},"<="&$C$4)', INT)
r += 1
row_new_units_month = r
metric_row(r, "New Units Added This Month (Management Agreement Date)",
           f'=COUNTIFS({col(PM,"A")},">="&DATE(YEAR($C$4),MONTH($C$4),1),{col(PM,"A")},"<="&$C$4)', INT)
r += 1
row_new_units_quarter = r
metric_row(r, "New Units Added This Quarter (Management Agreement Date)",
           f'=COUNTIFS({col(PM,"A")},">="&{qstart},{col(PM,"A")},"<="&$C$4)', INT)
r += 2

# =========================================================
# SECTION 5 — Data Health
# =========================================================
merge_banner(r, "5.  Data Health (needs manual cleanup)", section_font, section_fill, height=22)
r += 1
table_header(r, ["Metric", "Value", "Notes"])
r += 1
row_invalid_mgmt_date_pm = r
metric_row(r, 'PM ONLY — Invalid/Non-Date "Management Agreement Date" entries',
           f'=SUMPRODUCT(({col(PM,"A")}<>"")*(ISNUMBER({col(PM,"A")})=FALSE))', INT,
           'E.g. literal "???" placeholders. Excluded automatically from all date-based KPIs above; still needs manual correction at the source.')
r += 1
row_invalid_movein_pm = r
metric_row(r, 'PM ONLY — Invalid/Non-Date "Lease Move In" entries',
           f'=SUMPRODUCT(({col(PM,"M")}<>"")*(ISNUMBER({col(PM,"M")})=FALSE))', INT,
           'E.g. a free-text date such as "Sept 01, 2026" instead of a real date value.')
r += 1
row_invalid_mgmt_date_lo = r
metric_row(r, 'LEASE ONLY — Invalid/Non-Date "Management Agreement Date" entries',
           f'=SUMPRODUCT(({col(LO,"A")}<>"")*(ISNUMBER({col(LO,"A")})=FALSE))', INT,
           'E.g. a mistyped year such as "February 19, 0206".')
r += 1
row_missing_tenant_name = r
metric_row(r, "Occupied Units Missing Tenant Name",
           f'=SUMPRODUCT(({col(PM,"H")}="Occupied")*({col(PM,"J")}=""))', INT)
r += 1
row_missing_email = r
metric_row(r, "Occupied Units Missing Email",
           f'=SUMPRODUCT(({col(PM,"H")}="Occupied")*({col(PM,"K")}=""))', INT)
r += 1
row_missing_phone = r
metric_row(r, "Occupied Units Missing Phone Number",
           f'=SUMPRODUCT(({col(PM,"H")}="Occupied")*({col(PM,"L")}=""))', INT)
r += 2

# =========================================================
# Assumptions / notes footer
# =========================================================
merge_banner(r, "Assumptions & Data-Quality Notes", subheader_font, subheader_fill, height=18)
r += 1
notes = [
    "All KPIs pull live from PM ONLY, LEASE ONLY, and Len-Onboarded via SUMIFS/COUNTIFS/SUMPRODUCT — nothing on this tab is a hardcoded total.",
    "Formula ranges extend well past the current last data row (e.g. PM ONLY 2:2000) so new rows added to the source tabs are picked up automatically without editing this tab.",
    "COUNTIFS/SUMIFS date-range criteria (\">=\"/\"<=\") only match real date/number values, so non-date text like \"???\" is silently skipped rather than causing a formula error — see the Data Health section for a count of what still needs manual cleanup.",
    "\"TOTAL FEES WITH GST\" in PM ONLY has no values in the source file, so GST-inclusive management-fee revenue is estimated on this tab using the GST Rate input cell (C5, default 5%), matching the 5% GST convention already used on the LEASE ONLY tab's own formulas.",
    "\"Assigned Leasing Agent\" is empty across every PM ONLY row in the source data, so it is omitted from the workload breakdown.",
    "PM ONLY, LEASE ONLY, TENANTS, REALTOR FEES, SENT PM AGREEMENTS, and Len-Onboarded are unchanged from the uploaded file — only this Dashboard tab was added.",
]
for note in notes:
    ws.merge_cells(f"B{r}:F{r}")
    set_cell(f"B{r}", f"•  {note}", font=note_font, align=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[r].height = 26
    r += 1

# =========================================================
# Named ranges — stable handles for the Google Sheets API /
# Apps Script (survive Dashboard layout edits; a defined name
# is a level of indirection an API consumer can rely on
# instead of hardcoding cell coordinates).
# =========================================================
from openpyxl.workbook.defined_name import DefinedName

named_cells = {
    "KPI_ReportDate": 4,
    "KPI_GSTRate": 5,
    "KPI_TotalUnits": row_total_units,
    "KPI_ActiveUnits": row_active,
    "KPI_TerminatedUnits": row_terminated,
    "KPI_OccupiedUnits": row_occupied,
    "KPI_VacantUnits": row_vacant,
    "KPI_TakeOverUnits": row_takeover,
    "KPI_OccupancyRate": row_occ_rate,
    "KPI_RentRoll_PMOnly": 31,
    "KPI_RentRoll_LeaseOnly": 32,
    "KPI_MgmtFeeRevenue_ExclGST": row_mgmt_fee,
    "KPI_MgmtFeeRevenue_InclGST_Est": row_mgmt_fee + 1,
    "KPI_LeasingFeeRevenue_ExclGST_PMOnly": row_mgmt_fee + 2,
    "KPI_LeasingFeeRevenue_InclGST_LeaseOnly": row_mgmt_fee + 3,
    "KPI_OwnerDisbursements_PMOnly": row_disb_pm,
    "KPI_OwnerDisbursements_LeaseOnly": row_disb_lo,
    "KPI_OwnerDisbursements_Combined": row_disb_lo + 1,
    "KPI_AncillaryIncome_PetFee": row_disb_lo + 2,
    "KPI_AncillaryIncome_Parking": row_disb_lo + 3,
    "KPI_AncillaryIncome_Storage": row_disb_lo + 4,
    "KPI_LeasesExpiring30d": row_expiring[30],
    "KPI_LeasesExpiring60d": row_expiring[60],
    "KPI_LeasesExpiring90d": row_expiring[90],
    "KPI_VacantActiveUnits": row_vacant_active,
    "KPI_NewLeasesThisMonth": row_new_leases_month,
    "KPI_NewOwnersThisMonth": row_new_owners_month,
    "KPI_NewOwnersThisQuarter": row_new_owners_quarter,
    "KPI_NewUnitsThisMonth": row_new_units_month,
    "KPI_NewUnitsThisQuarter": row_new_units_quarter,
    "KPI_InvalidMgmtAgreementDate_PMOnly": row_invalid_mgmt_date_pm,
    "KPI_InvalidLeaseMoveIn_PMOnly": row_invalid_movein_pm,
    "KPI_InvalidMgmtAgreementDate_LeaseOnly": row_invalid_mgmt_date_lo,
    "KPI_OccupiedMissingTenantName": row_missing_tenant_name,
    "KPI_OccupiedMissingEmail": row_missing_email,
    "KPI_OccupiedMissingPhone": row_missing_phone,
}

# Sanity-check every hardcoded/offset row against its actual label before
# naming it, so a layout change trips a loud error instead of silently
# naming the wrong cell.
expected_labels = {
    31: "Total Monthly Rent Roll (PM ONLY)",
    32: "Total Monthly Rent Roll (LEASE ONLY, supplemental)",
    row_mgmt_fee + 1: "Total Management Fee Revenue, incl. GST (estimated)",
    row_mgmt_fee + 2: "Total Leasing Fee Revenue, excl. GST (PM ONLY)",
    row_mgmt_fee + 3: "Total Leasing Fee Revenue, incl. GST (LEASE ONLY)",
    row_disb_lo + 1: "Combined Total Owner Disbursements",
    row_disb_lo + 2: "Ancillary Income — Pet Fee (PM ONLY)",
    row_disb_lo + 3: "Ancillary Income — Parking (PM ONLY)",
    row_disb_lo + 4: "Ancillary Income — Storage (PM ONLY)",
}
for row_num, expected in expected_labels.items():
    actual = ws.cell(row=row_num, column=2).value
    assert actual == expected, f"Row {row_num} label mismatch: expected {expected!r}, got {actual!r}"

# Drop any named ranges from a previous build before re-adding them.
for existing_name in list(wb.defined_names.keys()):
    if existing_name.startswith("KPI_") or existing_name == "PM_Workload_Table":
        del wb.defined_names[existing_name]

for name, row_num in named_cells.items():
    wb.defined_names[name] = DefinedName(name, attr_text=f"'Dashboard'!$C${row_num}")

# Whole PM-workload table (header row + all PM rows) as one named range,
# for API consumers that want the full breakdown in one range read.
wb.defined_names["PM_Workload_Table"] = DefinedName(
    "PM_Workload_Table", attr_text=f"'Dashboard'!$B${pm_table_start - 1}:$E${pm_table_end}"
)

wb.save(SRC)
print("Dashboard sheet written. Last row used:", r)
print("PM table rows:", pm_table_start, "-", pm_table_end)
print("Named ranges defined:", len(named_cells) + 1)
